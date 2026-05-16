"""Local document loading and text extraction."""

from __future__ import annotations

import csv
import hashlib
import json
from abc import ABC, abstractmethod
from pathlib import Path

from bs4 import BeautifulSoup

from rag_mcp.config import Settings
from rag_mcp.schemas import Document


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, file_path: Path) -> list[Document]:
        """Extract one or more text documents from a file."""


class TextFileParser(DocumentParser):
    def parse(self, file_path: Path) -> list[Document]:
        content = self._read_text(file_path)
        return [_build_document(file_path, content, {"extension": file_path.suffix.lower()})]

    def _read_text(self, file_path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return file_path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return file_path.read_text(encoding="utf-8", errors="replace")


class HtmlParser(TextFileParser):
    def parse(self, file_path: Path) -> list[Document]:
        html = self._read_text(file_path)
        soup = BeautifulSoup(html, "html.parser")
        title = soup.title.string.strip() if soup.title and soup.title.string else file_path.stem
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        content = soup.get_text("\n")
        metadata = {"extension": file_path.suffix.lower(), "html_title": title}
        return [_build_document(file_path, content, metadata, title=title)]


class PdfParser(DocumentParser):
    def parse(self, file_path: Path) -> list[Document]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("缺少 pypdf 依赖，无法解析 PDF 文件。") from exc

        reader = PdfReader(str(file_path))
        documents: list[Document] = []
        for index, page in enumerate(reader.pages, start=1):
            content = page.extract_text() or ""
            if content.strip():
                documents.append(
                    _build_document(
                        file_path,
                        content,
                        {"extension": ".pdf", "page_no": index},
                        title=f"{file_path.stem} p{index}",
                        identity_suffix=f"page:{index}",
                    )
                )
        return documents


class DocxParser(DocumentParser):
    def parse(self, file_path: Path) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx 依赖，无法解析 DOCX 文件。") from exc

        doc = DocxDocument(str(file_path))
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        table_lines: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_lines.append(" | ".join(cells))
        content = "\n\n".join(paragraphs + table_lines)
        return [_build_document(file_path, content, {"extension": ".docx"})]


class CsvParser(DocumentParser):
    def parse(self, file_path: Path) -> list[Document]:
        rows: list[str] = []
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames:
                for row_number, row in enumerate(reader, start=1):
                    parts = [f"{key}: {value}" for key, value in row.items() if value not in (None, "")]
                    if parts:
                        rows.append(f"row {row_number}\n" + "\n".join(parts))
            else:
                handle.seek(0)
                plain_reader = csv.reader(handle)
                for row_number, row in enumerate(plain_reader, start=1):
                    if any(cell.strip() for cell in row):
                        rows.append(f"row {row_number}: " + " | ".join(row))
        return [_build_document(file_path, "\n\n".join(rows), {"extension": ".csv"})]


class ExcelParser(DocumentParser):
    def parse(self, file_path: Path) -> list[Document]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("缺少 openpyxl 依赖，无法解析 Excel 文件。") from exc

        workbook = load_workbook(str(file_path), read_only=True, data_only=True)
        documents: list[Document] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value) if value is not None else "" for value in next(iterator, [])]
            for row_number, row in enumerate(iterator, start=2):
                values = ["" if value is None else str(value) for value in row]
                if not any(values):
                    continue
                if any(headers):
                    parts = [
                        f"{header or f'column_{index + 1}'}: {value}"
                        for index, value in enumerate(values)
                        if value
                    ]
                    rows.append(f"row {row_number}\n" + "\n".join(parts))
                else:
                    rows.append(f"row {row_number}: " + " | ".join(values))
            if rows:
                documents.append(
                    _build_document(
                        file_path,
                        "\n\n".join(rows),
                        {"extension": ".xlsx", "sheet_name": sheet.title},
                        title=f"{file_path.stem} - {sheet.title}",
                        identity_suffix=f"sheet:{sheet.title}",
                    )
                )
        workbook.close()
        return documents


class JsonParser(TextFileParser):
    def parse(self, file_path: Path) -> list[Document]:
        raw_text = self._read_text(file_path)
        if file_path.suffix.lower() == ".jsonl":
            lines = [self._format_json(json.loads(line)) for line in raw_text.splitlines() if line.strip()]
            content = "\n\n".join(lines)
        else:
            content = self._format_json(json.loads(raw_text))
        return [_build_document(file_path, content, {"extension": file_path.suffix.lower()})]

    def _format_json(self, value) -> str:
        return json.dumps(value, ensure_ascii=False, indent=2)


class DocumentLoader:
    """Load local files or directories into normalized text documents."""

    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        text_parser = TextFileParser()
        self.parsers: dict[str, DocumentParser] = {
            ".txt": text_parser,
            ".md": text_parser,
            ".markdown": text_parser,
            ".html": HtmlParser(),
            ".htm": HtmlParser(),
            ".pdf": PdfParser(),
            ".docx": DocxParser(),
            ".csv": CsvParser(),
            ".xlsx": ExcelParser(),
            ".json": JsonParser(),
            ".jsonl": JsonParser(),
        }

    def load_path(self, path: str, recursive: bool = True) -> list[Document]:
        root = Path(path).expanduser().resolve()
        if not root.exists():
            raise FileNotFoundError(f"路径不存在：{root}")

        if root.is_file():
            return self._load_file(root)

        documents: list[Document] = []
        for file_path in self._iter_files(root, recursive):
            documents.extend(self._load_file(file_path))
        return documents

    def _iter_files(self, root: Path, recursive: bool) -> list[Path]:
        pattern = "**/*" if recursive else "*"
        files = [
            path
            for path in root.glob(pattern)
            if path.is_file() and path.suffix.lower() in self.settings.supported_extensions
        ]
        return sorted(files)

    def _load_file(self, file_path: Path) -> list[Document]:
        extension = file_path.suffix.lower()
        if extension not in self.settings.supported_extensions:
            raise ValueError(f"暂不支持的文件类型：{extension}")

        size = file_path.stat().st_size
        if size > self.settings.max_file_size_bytes:
            raise ValueError(f"文件过大：{file_path}，大小 {size} bytes，上限 {self.settings.max_file_size_bytes} bytes。")

        parser = self.parsers.get(extension)
        if parser is None:
            raise ValueError(f"没有可用解析器处理文件类型：{extension}")

        documents = parser.parse(file_path)
        for document in documents:
            document.metadata.setdefault("size_bytes", size)
            document.metadata.setdefault("source_path", str(file_path))
        return documents


def _build_document(
    file_path: Path,
    content: str,
    metadata: dict[str, object],
    title: str | None = None,
    identity_suffix: str = "",
) -> Document:
    document_id = _build_document_id(file_path, content, identity_suffix)
    return Document(
        document_id=document_id,
        title=title or file_path.stem,
        source_path=str(file_path),
        content=content,
        metadata={
            **metadata,
            "source_path": str(file_path),
        },
    )


def _build_document_id(file_path: Path, content: str, identity_suffix: str = "") -> str:
    digest = hashlib.sha256()
    digest.update(str(file_path).encode("utf-8"))
    digest.update(b"\0")
    digest.update(identity_suffix.encode("utf-8"))
    digest.update(b"\0")
    digest.update(content.encode("utf-8"))
    return digest.hexdigest()[:24]
